"""Extrae encabezados y catálogos institucionales; no copia expedientes ni nombres."""
import json
from pathlib import Path
import openpyxl

root=Path(__file__).resolve().parents[2]
back=root/'casos-app-backend/src/casos-conflicto'
front=root/'casos-app/src/app/casos/conflicto-gestion'
front.mkdir(exist_ok=True)
w=openpyxl.load_workbook(r'F:/SERVIDOR/BASES 2026/3. BDD CONFLICTO/DEIC53-2026.xlsx',read_only=True,data_only=True)
s=w['DEIC53']
headers=list(next(s.iter_rows(max_row=1,max_col=74,values_only=True)))
(back/'conflicto-excel-encabezados.json').write_text(json.dumps(headers,ensure_ascii=False,indent=2),encoding='utf-8')
fields=[
 ('seccionConoce','Sección que conoce','text','Recepción'),
 ('hechoInvestigar','Hecho a investigar','text','Recepción'),
 ('fechaAsignacion','Fecha de asignación al investigador','date','Recepción'),
 ('institucionRemitente','Institución remitente','text','Recepción'),
 ('procedencia','Procedencia','text','Recepción'),
 ('tipoDocumento','Tipo de documento','text','Recepción'),
 ('fechaDocumento','Fecha del documento','date','Recepción'),
 ('emplazamiento','Emplazamiento','text','Recepción'),
 ('firmanteNombre','Nombre de quien firma el requerimiento','text','Recepción'),
 ('firmanteCargo','Cargo de quien firma el requerimiento','text','Recepción'),
 ('lineamientosMp','Requerimientos y lineamientos del MP','textarea','Recepción'),
 ('resenaHecho','Reseña del hecho','textarea','Hechos'),
 ('movilHecho','Móvil del hecho','text','Hechos'),
 ('modusOperandi','Modus operandi','textarea','Hechos'),
 ('hipotesisHecho','Hipótesis del hecho','textarea','Hechos'),
 ('fechaHecho','Fecha del hecho','date','Hechos'),
 ('horaHecho','Hora del hecho','time','Hechos'),
 ('lugarHecho','Lugar del hecho','text','Hechos'),
 ('lugarDetencion','Lugar de la detención','textarea','Hechos'),
 ('aldeaColoniaBarrio','Aldea / colonia / barrio','text','Hechos'),
 ('zonaHecho','Zona','text','Hechos'),
 ('demarcacionComisaria','Demarcación de comisaría','text','Hechos'),
 ('tipoVehiculo','Tipo de vehículo','text','Vehículo'),
 ('tipoPlacas','Tipo de placas','text','Vehículo'),
 ('numeroPlacas','Número de placas','text','Vehículo'),
 ('tipoAtencion','Turno / FNNA','text','Recepción'),
 ('fechaPrimerInforme','Fecha del primer informe','date','Informe y entrega'),
 ('folios','Folios','number','Informe y entrega'),
 ('entrevistas','Entrevistas realizadas','number','Informe y entrega'),
 ('vistoBueno','Visto bueno físico (referencia)','text','Informe y entrega'),
 ('clasificacionCaso','Clasificación del caso','text','Informe y entrega'),
 ('fechaEntregaAnalisis','Fecha de entrega a análisis','date','Informe y entrega'),
 ('observaciones','Observaciones','textarea','Informe y entrega'),
 ('columnaAdicional','Dato adicional de la base (Columna1)','text','Informe y entrega'),
]
dto="import { IsOptional, IsString, IsDateString, IsInt, Min, Max, MaxLength, Matches } from 'class-validator';\nexport class FichaConflictoDto {\n"
schema="import { Prop, Schema, SchemaFactory } from '@nestjs/mongoose';\n@Schema({_id:false})\nexport class FichaConflicto {\n"
for k,l,t,section in fields:
 typ='number' if t=='number' else 'string'
 validation='@IsInt() @Min(0) @Max(1000000)' if t=='number' else '@IsDateString()' if t=='date' else '@Matches(/^([01]\\d|2[0-3]):[0-5]\\d$/)' if t=='time' else '@IsString() @MaxLength(10000)'
 dto+=f'  @IsOptional() {validation} {k}?: {typ};\n'
 options='type:Number,min:0,max:1000000,validate:Number.isInteger' if t=='number' else 'type:Date' if t=='date' else 'type:String,trim:true'
 schema+=f'  @Prop({{{options}}}) {k}?: {"Date" if t=="date" else typ};\n'
(back/'dto/ficha-conflicto.dto.ts').write_text(dto+'}\n',encoding='utf-8')
(back/'entities/ficha-conflicto.schema.ts').write_text(schema+'}\nexport const FichaConflictoSchema=SchemaFactory.createForClass(FichaConflicto);\n',encoding='utf-8')
(front/'conflicto-campos.ts').write_text('export const CAMPOS_CONFLICTO: {key:string;label:string;type:string;section:string}[] = '+json.dumps([dict(key=k,label=l,type=t,section=section) for k,l,t,section in fields],ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
cols={'seccionConoce':4,'institucionRemitente':8,'procedencia':9,'tipoDocumento':10,'tipoAtencion':66,'tipoVehiculo':63,'tipoPlacas':64,'sexo':42}
cats={k:set() for k in cols}
for row in s.iter_rows(min_row=2,max_col=74,values_only=True):
 for k,c in cols.items():
  if row[c] is not None: cats[k].add(str(row[c]).strip())
opts={k:sorted(v) for k,v in cats.items()}
opts['clasificacionCaso']=['A','B','C','D']
(front/'conflicto-catalogos.ts').write_text('// Etiquetas institucionales y categorías de la referencia; no contiene personas.\nexport const OPCIONES_CONFLICTO:Record<string,readonly string[]> = '+json.dumps(opts,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
